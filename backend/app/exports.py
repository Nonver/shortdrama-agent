"""剧本导出：Markdown / JSON / Word / Excel(分镜表) / PDF。"""
from __future__ import annotations
import io
import json

from .generator import render_markdown


def to_markdown(script: dict) -> str:
    return render_markdown(script)


def to_json(script: dict) -> bytes:
    return json.dumps(script, ensure_ascii=False, indent=2).encode("utf-8")


def to_docx(script: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    p = script["project"]
    doc.add_heading(p["项目名"], level=0)
    doc.add_heading("一、工程信息", level=1)
    for k, v in p.items():
        doc.add_paragraph(f"{k}：{v}", style="List Bullet")
    for ep in script["episodes"]:
        doc.add_heading(f"第 {ep['集号']} 集 ｜ {ep['本集主线']}", level=1)
        for sc in ep["scenes"]:
            doc.add_heading(f"场 {sc['场号']} ｜ {sc['场景']} ｜ 功能：{sc['本场功能']}", level=2)
            for sh in sc["shots"]:
                doc.add_paragraph(f"镜 {sh['镜号']} 〔{sh['景别']}·{sh['机位角度']}·{sh['运镜']}·{sh['构图']}〕 {sh['镜头标记']}")
                doc.add_paragraph(f"画面：{sh['画面描述']}", style="List Bullet")
                if sh["台词"]:
                    doc.add_paragraph(f"台词：{sh['台词']}", style="List Bullet")
                if sh["OS旁白/内心独白"]:
                    doc.add_paragraph(f"旁白/独白：{sh['OS旁白/内心独白']}", style="List Bullet")
                if sh["人物心理活动"]:
                    doc.add_paragraph(f"心理：{sh['人物心理活动']}", style="List Bullet")
                doc.add_paragraph(f"服化道：{sh['服化道']} ｜ 色调：{sh['色调']} ｜ 灯光：{sh['灯光']}", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_xlsx(script: dict) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"
    headers = ["集", "场", "镜号", "景别", "机位角度", "运镜", "构图", "镜头标记",
               "画面描述", "台词", "OS旁白/内心独白", "人物心理活动", "服化道",
               "色调", "灯光", "BGM", "音效", "字幕花字", "单镜时长", "情绪强度", "节拍", "标记", "负责Agent"]
    ws.append(headers)
    for ep in script["episodes"]:
        for sc in ep["scenes"]:
            for sh in sc["shots"]:
                ws.append([ep["集号"], sc["场号"], sh["镜号"], sh["景别"], sh["机位角度"],
                           sh["运镜"], sh["构图"], sh["镜头标记"], sh["画面描述"],
                           sh["台词"], sh["OS旁白/内心独白"], sh["人物心理活动"],
                           sh["服化道"], sh["色调"], sh["灯光"], sh["BGM"], sh["音效"],
                           sh["字幕花字"], sh["单镜时长"], sh["情绪强度"], sh["节拍"],
                           sh["标记"], sh["负责Agent"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(script: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import mm

    styles = getSampleStyleSheet()
    cn = ParagraphStyle("cn", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12)

    flow = [Paragraph(script["project"]["项目名"], h1)]
    for ep in script["episodes"]:
        flow.append(Paragraph(f"第 {ep['集号']} 集 ｜ {ep['本集主线']}", h2))
        for sc in ep["scenes"]:
            flow.append(Paragraph(f"场 {sc['场号']} ｜ {sc['场景']} ｜ 功能：{sc['本场功能']}", cn))
            for sh in sc["shots"]:
                line = (f"镜{sh['镜号']}〔{sh['景别']}·{sh['运镜']}〕{sh['镜头标记']}："
                        f"{sh['画面描述']} {sh['台词']} {sh['OS旁白/内心独白']} "
                        f"色调{sh['色调']}/灯光{sh['灯光']} 时长{sh['单镜时长']} {sh['标记']}")
                flow.append(Paragraph(line, cn))
            flow.append(Spacer(1, 4))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            title=script["project"]["项目名"])
    doc.build(flow)
    return buf.getvalue()
